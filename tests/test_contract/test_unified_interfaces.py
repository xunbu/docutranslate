"""Cross-entry regressions for every supported workflow, without external APIs."""
import base64
import inspect
import json
import logging
from dataclasses import fields
from pathlib import Path

import httpx
import pytest
from pydantic import TypeAdapter

from docutranslate.core.factory import create_workflow_from_payload
from docutranslate.core.file_types import INPUT_WORKFLOWS, WORKFLOW_FORMATS, detect_workflow
from docutranslate.core.params import resolve_auto_payload, resolve_payload
from docutranslate.core.schemas import BaseWorkflowParams, TranslatePayload
from docutranslate.sdk import Client
from docutranslate.server.core import TranslationService
from test_mcp_workflow_sdk_fixes import fake_mcp_server_module, FakeService


@pytest.mark.parametrize("workflow", WORKFLOW_FORMATS)
def test_every_workflow_shares_config_exports_and_progress(workflow):
    payload = resolve_payload({"workflow_type": workflow, "skip_translate": True,
                               "convert_engine": "identity", "md2docx_engine": "python",
                               "streaming": False, "stream_idle_timeout": 17,
                               "non_stream_timeout": None}, f"sample.{workflow}")
    progress = object()
    direct = create_workflow_from_payload(payload)
    server = TranslationService()._create_workflow(
        payload, logging.getLogger("test"), progress, lambda: None, "python")
    for field in fields(direct.config.translator_config):
        if field.name not in {"logger", "progress_tracker"}:
            assert getattr(direct.config.translator_config, field.name) == getattr(server.config.translator_config, field.name)
    assert server.config.progress_tracker is progress
    assert server.config.translator_config.progress_tracker is progress
    assert set(TranslationService()._build_export_map(server, "sample", True)) == set(WORKFLOW_FORMATS[workflow])


@pytest.mark.parametrize("suffix", INPUT_WORKFLOWS)
def test_routing_registry_is_shared(suffix):
    assert Client()._detect_workflow(Path("sample" + suffix)) == detect_workflow("sample" + suffix)


@pytest.mark.parametrize("suffix", [".doc", ".xls", ".ppt"])
def test_legacy_office_files_rejected_before_translation(suffix):
    with pytest.raises(ValueError, match="请先转换"):
        resolve_payload({"workflow_type": "auto", "skip_translate": True}, "sample" + suffix)
    assert suffix.lstrip(".") not in Client().capabilities["input_formats"]


def test_auto_preserves_empty_separator_and_disabled_docx():
    auto = TypeAdapter(TranslatePayload).validate_python({
        "workflow_type": "auto", "skip_translate": True, "separator": "",
        "md2docx_engine": None, "streaming": False, "non_stream_timeout": None})
    assert resolve_auto_payload(auto, "sample.txt").separator == ""
    markdown = resolve_auto_payload(auto, "sample.md")
    assert markdown.md2docx_engine is None
    assert markdown.streaming is False
    assert markdown.non_stream_timeout is None


def test_auto_inherits_target_parser_defaults(monkeypatch):
    import docutranslate.core.schemas as schemas
    monkeypatch.setattr(schemas.MarkdownWorkflowParams.model_fields["mineru_deploy_image_analysis"], "default", False)
    # Rebuilding the target model makes its changed default observable by validation.
    schemas.MarkdownWorkflowParams.model_rebuild(force=True)
    try:
        auto = TypeAdapter(TranslatePayload).validate_python({"workflow_type": "auto", "skip_translate": True})
        assert resolve_auto_payload(auto, "sample.md").mineru_deploy_image_analysis is False
    finally:
        monkeypatch.undo()
        schemas.MarkdownWorkflowParams.model_rebuild(force=True)


def test_zero_false_and_empty_prompt_override_environment(monkeypatch):
    import docutranslate.core.schemas as schemas
    monkeypatch.setattr(schemas, "ENV_FORCE_OVERRIDE", False)
    monkeypatch.setattr(schemas, "TEMPERATURE", 0.8)
    monkeypatch.setattr(schemas, "SYSTEM_PROXY_ENABLE", True)
    monkeypatch.setattr(schemas, "CUSTOM_PROMPT", "default prompt")
    value = BaseWorkflowParams(base_url="https://example.invalid", model_id="test",
                               temperature=0, system_proxy_enable=False, custom_prompt="")
    assert (value.temperature, value.system_proxy_enable, value.custom_prompt) == (0, False, "")


@pytest.mark.asyncio
async def test_mcp_overrides_and_public_parameter_coverage(fake_mcp_server_module, tmp_path):
    service = FakeService()
    server = fake_mcp_server_module.create_mcp_server(translation_service=service)
    submit = server.tools["submit_task"]
    source = tmp_path / "sample.txt"
    source.write_text("hello", encoding="utf-8")
    await server.tools["configure_client"](custom_prompt="default", separator="/", non_stream_timeout=30)
    await submit(str(source), workflow_type="txt", skip_translate=True, separator="",
                 custom_prompt="", temperature=0, streaming=False, connect_timeout=19,
                 non_stream_timeout=None, read_timeout=None)
    payload = service.payloads[-1]
    assert payload.separator == payload.custom_prompt == ""
    assert payload.temperature == 0 and payload.streaming is False
    assert payload.connect_timeout == 19
    assert payload.non_stream_timeout is None and payload.read_timeout is None
    public = set(BaseWorkflowParams.model_fields)
    exposed = set(inspect.signature(submit).parameters)
    aliases = {"extra_body": "extra_body_json", "glossary_dict": "glossary_dict_json",
               "glossary_agent_config": "glossary_agent_config_json"}
    assert {aliases.get(field, field) for field in public} <= exposed


@pytest.mark.asyncio
async def test_mcp_workflow_defaults_json_and_null_docx(fake_mcp_server_module, tmp_path):
    service = FakeService()
    server = fake_mcp_server_module.create_mcp_server(translation_service=service)
    source = tmp_path / "sample.md"
    source.write_text("hello", encoding="utf-8")
    submit = server.tools["submit_task"]
    await submit(str(source), workflow_type="html", skip_translate=True)
    assert service.payloads[-1].separator == " "
    await submit(str(source), workflow_type="json", skip_translate=True)
    assert service.payloads[-1].json_paths == ["$..*"]
    response = await submit(str(source), workflow_type="markdown_based", skip_translate=True, md2docx_engine=None)
    assert response.startswith("Translation task submitted"), response
    assert service.payloads[-1].md2docx_engine is None
    assert "html" in fake_mcp_server_module._get_formats_info()["output_formats_by_workflow"]["txt"]


@pytest.mark.asyncio
async def test_mcp_auto_pdf_leaves_parser_selection_to_shared_resolver(fake_mcp_server_module, monkeypatch, tmp_path):
    monkeypatch.setattr(fake_mcp_server_module, "CONVERT_ENGINE", "identity")
    service = FakeService()
    server = fake_mcp_server_module.create_mcp_server(translation_service=service)
    source = tmp_path / "sample.pdf"
    source.write_bytes(b"synthetic")
    await server.tools["submit_task"](str(source), skip_translate=True)
    assert service.payloads[-1].convert_engine is None
    assert resolve_auto_payload(service.payloads[-1], source.name).convert_engine == "mineru"


@pytest.mark.asyncio
async def test_mcp_state_release_logs_and_glossary_roundtrip(fake_mcp_server_module, tmp_path):
    from docutranslate.glossary.glossary import Glossary
    class Service(FakeService):
        def get_task_state(self, task_id):
            return {"cancelled": True, "statistics": {"requests": 2}} if task_id == "known" else None
        async def release_task(self, task_id):
            return {"released": False, "message": "missing"}
        def list_tasks(self):
            return ["known"]
        def get_task_logs(self, task_id):
            return ["retained log"]
    server = fake_mcp_server_module.create_mcp_server(translation_service=Service())
    status = await server.tools["get_task_status"]("known")
    assert '"cancelled"' in status and '"requests": 2' in status
    assert (await server.tools["release_task"]("missing")).startswith("Error")
    assert "retained log" in await server.tools["get_task_logs"]("known")
    assert "known" in await server.tools["list_tasks"]()
    source = tmp_path / "glossary.csv"
    source.write_bytes(Glossary.glossary_dict2csv({"Hello": "你好"}).content)
    glossary = await server.tools["load_glossary_file"](str(source))
    assert "(1 entries)" in glossary and "src" not in glossary


@pytest.mark.asyncio
@pytest.mark.parametrize("endpoint", ["json", "file", "flat"])
async def test_all_http_inputs_forward_options(monkeypatch, endpoint):
    import docutranslate.app as module
    captured = []
    async def capture(**kwargs):
        captured.append(kwargs["payload"])
        return {"task_started": True, "task_id": kwargs["task_id"]}
    monkeypatch.setattr(module.translation_service, "start_translation", capture)
    monkeypatch.setattr(module.translation_service, "get_task_state", lambda _: {"download_ready": True})
    payload = {"workflow_type": "auto", "skip_translate": True, "separator": "",
               "md2docx_engine": None, "streaming": False, "connect_timeout": 19,
               "office_password": "test", "html_exporter_options": {"cdn": False}}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=module.app), base_url="http://test") as client:
        if endpoint == "json":
            response = await client.post("/service/translate", json={"file_name": "sample.md", "file_content": base64.b64encode(b"hello").decode(), "payload": payload})
        elif endpoint == "file":
            response = await client.post("/service/translate/file", files={"file": ("sample.md", b"hello")}, data={"payload": json.dumps(payload)})
        else:
            data = {**payload, "skip_translate": "true", "streaming": "false", "md2docx_engine": "none",
                    "html_exporter_options_json": json.dumps(payload["html_exporter_options"])}
            data.pop("html_exporter_options")
            response = await client.post("/service/flat-translate", files={"file": ("sample.md", b"hello")}, data=data)
    assert response.status_code == 200, response.text
    actual = captured[-1]
    assert actual.separator == "" and actual.md2docx_engine is None
    assert actual.streaming is False and actual.connect_timeout == 19
    assert actual.office_password == "test"
    assert actual.html_exporter_options.cdn is False


@pytest.mark.asyncio
async def test_pptx_html_export_options_survive_service_export(tmp_path):
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Hidden audit slide"
    slide._element.set("show", "0")
    source = tmp_path / "sample.pptx"
    deck.save(str(source))
    result = await Client().translate_async(str(source), skip_translate=True,
                                           html_exporter_options={"cdn": False, "include_hidden_slides": True})
    assert b"Hidden audit slide" in base64.b64decode(result.export("html"))
    exports = TranslationService()._build_export_map(result.workflow, "sample", True)
    assert "Hidden audit slide" in exports["html"][0]()
    assert result.statistics == result.workflow.get_statistics()
    assert result.attachments is result.workflow.get_attachment()


@pytest.mark.asyncio
async def test_sdk_inherits_constructor_controls_and_can_clear_nullable_timeouts(tmp_path):
    source = tmp_path / "sample.txt"
    source.write_text("hello", encoding="utf-8")
    result = await Client(skip_translate=True, workflow_type="txt", read_timeout=17,
                          non_stream_timeout=29).translate_async(
        str(source), read_timeout=None, non_stream_timeout=None)
    config = result.workflow.config.translator_config
    assert config.skip_translate is True
    assert config.read_timeout is None and config.non_stream_timeout is None


@pytest.mark.asyncio
async def test_all_surfaces_share_capabilities_and_cancelled_statistics(fake_mcp_server_module, monkeypatch):
    import docutranslate.app as module
    expected = Client().capabilities
    state = {"is_processing": False, "cancelled": True, "error_flag": False,
             "download_ready": False, "status_message": "cancelled", "original_filename_stem": "sample",
             "task_start_time": 1, "task_end_time": 2, "statistics": {"requests": 3}}
    monkeypatch.setattr(module.translation_service, "get_task_state", lambda _: state)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=module.app), base_url="http://test") as client:
        assert (await client.get("/service/capabilities")).json() == expected
        status = (await client.get("/service/status/known")).json()
    mcp_status = fake_mcp_server_module._task_state_to_dict("known", state)
    assert status["status"] == mcp_status["status"] == "cancelled"
    assert status["statistics"] == mcp_status["statistics"] == {"requests": 3}
    assert fake_mcp_server_module._get_formats_info() == expected


@pytest.mark.asyncio
async def test_flat_translate_empty_select_fields_mean_not_provided(monkeypatch):
    from docutranslate.config import MD2DOCX_ENGINE, NON_STREAM_TIMEOUT, READ_TIMEOUT

    import docutranslate.app as module

    captured = []

    async def capture(**kwargs):
        captured.append(kwargs["payload"])
        return {"task_started": True, "task_id": kwargs["task_id"]}

    monkeypatch.setattr(module.translation_service, "start_translation", capture)
    monkeypatch.setattr(module.translation_service, "get_task_state",
                        lambda _: {"download_ready": True})

    # Legacy form clients send every optional select as an empty string. Only
    # separator/custom_prompt keep an empty value; everything else is unset.
    data = {
        "workflow_type": "auto", "skip_translate": "true",
        "base_url": "https://api.invalid/v1", "api_key": "test-api-key",
        "model_id": "gpt-test", "to_lang": "中文",
        "provider": "", "thinking": "", "convert_engine": "", "insert_mode": "",
        "segment_mode": "", "model_version": "", "mineru_language": "",
        "mineru_deploy_backend": "", "mineru_token": "", "office_password": "",
        "extra_body_json": "", "html_exporter_options_json": "",
        "separator": "", "custom_prompt": "", "md2docx_engine": "",
        "read_timeout": "", "non_stream_timeout": "",
    }
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=module.app),
                                 base_url="http://test") as client:
        response = await client.post(
            "/service/flat-translate",
            files={"file": ("sample.md", b"# Title\n")},
            data=data,
        )
    assert response.status_code == 200, response.text

    payload = captured[-1]
    assert payload.separator == "" and payload.custom_prompt == ""
    assert payload.provider is None
    assert getattr(payload, "office_password", None) is None
    resolved = resolve_auto_payload(payload, "sample.md")
    assert resolved.workflow_type == "markdown_based"
    # Empty md2docx_engine inherits the configured engine instead of disabling DOCX.
    assert resolved.md2docx_engine == MD2DOCX_ENGINE
    # The flat form cannot express JSON null: empty timeouts inherit defaults.
    assert resolved.read_timeout == READ_TIMEOUT
    assert resolved.non_stream_timeout == NON_STREAM_TIMEOUT


@pytest.mark.asyncio
async def test_mcp_empty_string_options_inherit_configuration(fake_mcp_server_module, tmp_path):
    from docutranslate.config import MD2DOCX_ENGINE

    service = FakeService()
    server = fake_mcp_server_module.create_mcp_server(translation_service=service)
    source = tmp_path / "sample.txt"
    source.write_text("hello", encoding="utf-8")
    markdown = tmp_path / "sample.md"
    markdown.write_text("# Title\n", encoding="utf-8")

    # Empty strings in configure_client must not poison the shared config.
    await server.tools["configure_client"](
        provider="", thinking="", insert_mode="", model_version="",
        separator="/", custom_prompt="default", md2docx_engine="",
    )
    response = await server.tools["submit_task"](
        str(source), workflow_type="txt", skip_translate=True,
        provider="", thinking="", convert_engine="", insert_mode="",
        separator="", custom_prompt="",
    )
    assert response.startswith("Translation task submitted"), response
    payload = service.payloads[-1]
    assert payload.provider is None
    assert payload.separator == "" and payload.custom_prompt == ""
    assert payload.insert_mode == "replace"  # empty enum means workflow default

    # Omitted options now inherit the configured values.
    await server.tools["submit_task"](str(source), workflow_type="txt", skip_translate=True)
    inherited = service.payloads[-1]
    assert inherited.separator == "/"
    assert inherited.custom_prompt == "default"

    # Empty md2docx_engine means inherit for auto payloads, JSON null disables it.
    await server.tools["submit_task"](str(markdown), skip_translate=True, md2docx_engine="")
    assert service.payloads[-1].md2docx_engine == MD2DOCX_ENGINE
    await server.tools["submit_task"](str(markdown), skip_translate=True, md2docx_engine=None)
    assert service.payloads[-1].md2docx_engine is None
