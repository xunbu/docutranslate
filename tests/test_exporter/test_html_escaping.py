"""Exporter hardening regressions.

* TXT/SRT/ASS text was injected into HTML unescaped (broken markup, live
  ``<script>``), and the SRT template read the misspelled ``tittle`` so every
  exported title was empty.
* The JSON viewer injected raw JSON into ``<script>`` (``</script>`` breakout)
  and pulled a second renderjson copy from a CDN with a malformed URL.
* ``PPTX2HTMLExporterConfig.include_hidden_slides`` was stored but never applied.
"""
import io

import pytest
from pptx import Presentation

from docutranslate.exporter.ass.ass2html_exporter import Ass2HTMLExporter
from docutranslate.exporter.js.json2html_exporter import Json2HTMLExporter
from docutranslate.exporter.pptx.pptx2html_exporter import (
    PPTX2HTMLExporter,
    PPTX2HTMLExporterConfig,
)
from docutranslate.exporter.srt.srt2html_exporter import Srt2HTMLExporter
from docutranslate.exporter.txt.txt2html_exporter import TXT2HTMLExporter
from docutranslate.ir.document import Document

PAYLOAD = "5 < 6 & a > b\n</pre><script>alert(1)</script>\n"


def test_txt_html_escapes_document_text():
    html = TXT2HTMLExporter().export(
        Document.from_bytes(content=PAYLOAD.encode(), suffix=".txt", stem="math")
    ).content.decode("utf-8")

    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html
    assert "&lt;" in html and "&amp;" in html


def test_srt_html_escapes_content_and_sets_title():
    srt = "1\n00:00:01,000 --> 00:00:02,000\nMy subtitle\na < b & c\n\n"
    html = Srt2HTMLExporter().export(
        Document.from_bytes(content=srt.encode(), suffix=".srt", stem="MySubtitle")
    ).content.decode("utf-8")

    assert "<title>MySubtitle</title>" in html
    assert "<title></title>" not in html
    assert "a &lt; b &amp; c" in html
    assert "<br>" in html


def test_ass_html_escapes_content():
    ass = "[Script Info]\nTitle: t\n\n[Events]\nFormat: Layer, Start, End, Style, Text\n"
    ass += "Dialogue: 0,0:00:00.00,0:00:02.00,Default,<b>&amp; not html</b>\n"
    html = Ass2HTMLExporter().export(
        Document.from_bytes(content=ass.encode(), suffix=".ass", stem="subs")
    ).content.decode("utf-8")

    assert "<b>&amp; not html</b>" not in html
    assert "&lt;b&gt;&amp;amp; not html&lt;/b&gt;" in html


def test_json_html_escapes_script_breakout():
    payload = '{"note": "</script><script>alert(1)</script>"}'
    html = Json2HTMLExporter().export(
        Document.from_bytes(content=payload.encode(), suffix=".json", stem="data")
    ).content.decode("utf-8")

    assert "</script><script>alert(1)</script>" not in html
    assert "\\u003c/script\\u003e" in html


def test_json_html_has_no_broken_cdn_script_tag():
    html = Json2HTMLExporter().export(
        Document.from_bytes(content=b"{}", suffix=".json", stem="data")
    ).content.decode("utf-8")

    assert "cdn.jsdelivr.net/npm/renderjson" not in html


def _pptx_with_hidden_slide() -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    first = presentation.slides.add_slide(layout)
    first.shapes.title.text = "Visible Title"
    hidden = presentation.slides.add_slide(layout)
    hidden.shapes.title.text = "Hidden Title"
    hidden.element.set("show", "0")

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_pptx_html_excludes_hidden_slides_by_default():
    content = _pptx_with_hidden_slide()
    html = PPTX2HTMLExporter().export(
        Document.from_bytes(content=content, suffix=".pptx", stem="deck")
    ).content.decode("utf-8")

    assert "Visible Title" in html
    assert "Hidden Title" not in html


def test_pptx_html_can_include_hidden_slides():
    content = _pptx_with_hidden_slide()
    exporter = PPTX2HTMLExporter(PPTX2HTMLExporterConfig(include_hidden_slides=True))
    html = exporter.export(
        Document.from_bytes(content=content, suffix=".pptx", stem="deck")
    ).content.decode("utf-8")

    assert "Hidden Title" in html


def test_passthrough_exporters_use_their_own_base_class():
    from docutranslate.exporter.docx.base import DocxExporter
    from docutranslate.exporter.epub.base import EpubExporter
    from docutranslate.exporter.epub.epub2epub_exporter import Epub2EpubExporter
    from docutranslate.exporter.js.base import JsonExporter
    from docutranslate.exporter.js.json2json_exporter import Json2JsonExporter
    from docutranslate.exporter.pptx.base import PPTXExporter
    from docutranslate.exporter.pptx.pptx2pptx_exporter import PPTX2PPTXExporter
    from docutranslate.exporter.xlsx.base import XlsxExporter

    assert isinstance(Epub2EpubExporter(), EpubExporter)
    assert not isinstance(Epub2EpubExporter(), XlsxExporter)
    assert isinstance(PPTX2PPTXExporter(), PPTXExporter)
    assert not isinstance(PPTX2PPTXExporter(), DocxExporter)
    assert isinstance(Json2JsonExporter(), JsonExporter)
